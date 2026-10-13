#!/usr/bin/env python3
"""Build the complete English FLOWER noise-mismatch report as a DOCX."""

from __future__ import annotations

import csv
import json
import math
import statistics
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "artifacts/experiments/main"
SUMMARY = RESULTS / "summary.csv"
PAIRED = RESULTS / "metrics_with_deltas.csv"
CONFIG = ROOT / "config/robustness.json"
FIGURES = ROOT / "artifacts/figures"
EQUATIONS = ROOT / "artifacts/build/report-equations"
OUTPUT = ROOT / "artifacts/report/flower-noise-scale-mismatch-report.docx"

NAVY = "17324D"
BLUE = "176B87"
PALE_BLUE = "EAF4F8"
MID_GRAY = "606A73"
LIGHT_BORDER = "D9D9D9"
BLACK = "000000"

DATASET_LABELS = {"celeba": "CelebA", "afhq_cat": "AFHQ-Cat"}
TASK_LABELS = {
    "denoising": "Denoising",
    "gaussian_deblurring": "Gaussian deblurring",
    "super_resolution": "Super-resolution",
    "random_inpainting": "Random inpainting",
    "box_inpainting": "Box inpainting",
}


def set_run_font(run, name: str = "Aptos", size: float | None = None, color: str = BLACK) -> None:
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), LIGHT_BORDER)


def set_cell_margins(cell, value: int = 95) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    properties.append(node)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size: float = 9.0,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat_header(table.rows[0])
    for cell, label, width in zip(table.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        run.bold = True
        set_run_font(run, size=font_size, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column_index in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            set_run_font(paragraph.add_run(value), size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_picture(doc: Document, path: Path, width: float, caption: str) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = doc.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(3)
    caption_paragraph.paragraph_format.space_after = Pt(9)
    set_run_font(caption_paragraph.add_run(caption), size=9, color=MID_GRAY)


def add_equation(doc: Document, filename: str, width: float) -> None:
    path = EQUATIONS / filename
    if not path.is_file():
        raise FileNotFoundError(f"Missing equation asset {path}; run tools/build_equation_assets.py")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.add_run().add_picture(str(path), width=Inches(width))


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run(text), size=10.4)


def add_reference(doc: Document, number: int, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(5)
    set_run_font(paragraph.add_run(f"[{number}] {text}"), size=9.2)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=8.5, color=MID_GRAY)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.4)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title_border = title._element.get_or_add_pPr().find(qn("w:pBdr"))
    if title_border is not None:
        title._element.get_or_add_pPr().remove(title_border)
    for style_name, size in (("Heading 1", 16.5), ("Heading 2", 12.5), ("Heading 3", 11.2)):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    caption = doc.styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    footer = section.footer
    footer.paragraphs[0].text = "FLOWER Under Measurement Noise Scale Mismatch"
    footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in footer.paragraphs[0].runs:
        set_run_font(run, size=8.3, color=MID_GRAY)
    add_page_number(footer.add_paragraph())


def load_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="") as stream:
        return list(csv.DictReader(stream))


def mean_abs_sensitivity(summary: list[dict[str, str]], dataset: str, task: str) -> float:
    values = [
        abs(float(row["psnr_delta_vs_correct_mean"]))
        for row in summary
        if row["dataset"] == dataset
        and row["task"] == task
        and float(row["assumed_to_true_sigma"]) != 1.0
    ]
    return statistics.mean(values)


def metric_with_ci(row: dict[str, str], metric: str, digits: int) -> str:
    return f"{float(row[f'{metric}_mean']):.{digits}f} +/- {float(row[f'{metric}_ci95']):.{digits}f}"


def validate_results(summary: list[dict[str, str]], paired: list[dict[str, str]]) -> None:
    if len(summary) != 50 or len(paired) != 1000:
        raise ValueError(f"Expected 50 summary rows and 1000 paired rows; found {len(summary)} and {len(paired)}")
    for row in summary:
        if int(row["n"]) != 20:
            raise ValueError("Every aggregate row must contain 20 images")
        for metric in ("psnr_mean", "ssim_mean", "lpips_mean"):
            if not math.isfinite(float(row[metric])):
                raise ValueError(f"Non-finite {metric} in {row['dataset']}/{row['task']}")


def main() -> None:
    config = json.loads(CONFIG.read_text())
    summary = load_rows(SUMMARY)
    paired = load_rows(PAIRED)
    validate_results(summary, paired)
    lookup = {
        (row["dataset"], row["task"], float(row["assumed_to_true_sigma"])): row
        for row in summary
    }
    setting_scores = {
        (dataset, task): mean_abs_sensitivity(summary, dataset, task)
        for dataset in DATASET_LABELS
        for task in TASK_LABELS
    }
    most_sensitive = max(setting_scores, key=setting_scores.get)
    least_sensitive = min(setting_scores, key=setting_scores.get)

    selected_dataset, selected_task = most_sensitive
    mismatch_rows = [
        row
        for row in summary
        if row["dataset"] == selected_dataset
        and row["task"] == selected_task
        and float(row["assumed_to_true_sigma"]) != 1.0
    ]
    worst_row = max(mismatch_rows, key=lambda row: abs(float(row["psnr_delta_vs_correct_mean"])))
    worst_ratio = float(worst_row["assumed_to_true_sigma"])
    setting_paired = [
        row
        for row in paired
        if row["dataset"] == selected_dataset
        and row["task"] == selected_task
        and float(row["assumed_to_true_sigma"]) == worst_ratio
    ]
    loss_fraction = sum(float(row["psnr_delta_vs_correct"]) < 0 for row in setting_paired) / len(setting_paired)

    doc = Document()
    configure_document(doc)

    accent = doc.add_paragraph()
    accent.paragraph_format.space_after = Pt(55)
    square = accent.add_run("■")
    set_run_font(square, size=15, color=BLUE)
    title = doc.add_paragraph(style="Title")
    title.add_run("FLOWER Under Measurement Noise Scale Mismatch")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(8)
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(
        subtitle.add_run("A paired robustness study across five linear inverse problems and two image domains"),
        name="Aptos Display",
        size=15,
        color=BLUE,
    )
    metadata = doc.add_paragraph()
    metadata.paragraph_format.space_before = Pt(26)
    set_run_font(metadata.add_run("Project report\nItay Shorian\n13 October 2026"), size=10.8, color=MID_GRAY)
    scope = doc.add_paragraph()
    scope.paragraph_format.space_before = Pt(140)
    set_run_font(
        scope.add_run("20 CelebA images  |  20 AFHQ-Cat images  |  1,000 reconstructions  |  PSNR SSIM LPIPS"),
        size=9.5,
        color=MID_GRAY,
    )
    doc.add_page_break()

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "FLOWER uses a pretrained flow-matching prior and a measurement-aware Gaussian refinement to solve linear inverse problems. "
        "The refinement requires a measurement-noise standard deviation, but real applications may not know it accurately. This study "
        "changes only the standard deviation supplied to FLOWER while holding each measurement, degradation operator, mask or kernel, "
        "model, and random trajectory fixed. The benchmark evaluates five ratios from 0.5 to 1.5 on 20 fixed images from each of the "
        "paper's CelebA and AFHQ-Cat 100-image test pools and covers denoising, deblurring, super-resolution, random inpainting, and box "
        "inpainting. The original model and solver remain unchanged. The largest mean absolute paired PSNR response occurs for "
        f"{DATASET_LABELS[selected_dataset]} {TASK_LABELS[selected_task].lower()} at {setting_scores[most_sensitive]:.2f} dB, while the "
        f"smallest occurs for {DATASET_LABELS[least_sensitive[0]]} {TASK_LABELS[least_sensitive[1]].lower()} at "
        f"{setting_scores[least_sensitive]:.2f} dB. The results therefore quantify where FLOWER needs an accurate noise scale and where "
        "moderate misspecification has limited effect."
    )

    doc.add_heading("Research question and contribution", level=1)
    doc.add_paragraph(
        "The research question is how reconstruction quality changes when the measurement was generated with the true noise scale but "
        "FLOWER receives a different scale. The project contributes a controlled paired benchmark, not a new generative model and not an "
        "additional noise estimator. Pairing removes measurement and solver randomness from comparisons between noise assumptions."
    )
    add_equation(doc, "measurement.png", 4.8)
    add_bullet(doc, "The covariance shape is fixed and correct in every experiment.")
    add_bullet(doc, "The assumed standard deviation ratios are 0.5, 0.75, 1.0, 1.25, and 1.5.")
    add_bullet(doc, "The corresponding variance ratios are 0.25, 0.5625, 1.0, 1.5625, and 2.25.")

    doc.add_heading("FLOWER method", level=1)
    doc.add_paragraph(
        "FLOWER starts from a source sample and advances along a straight-line flow. At each time it performs a prior prediction, a "
        "measurement refinement, and a progression step. The public experiments use the deterministic Flower1 setting with gamma equal "
        "to zero, so the refinement uses its conditional mean rather than drawing its covariance term [1]."
    )
    add_table(
        doc,
        ["Step", "Operation", "Dependence on assumed sigma"],
        [
            ["1", "Predict a flow-consistent destination from the velocity network", "No explicit dependence"],
            ["2", "Combine the destination with the measurement through a Gaussian refinement", "Direct dependence through inverse variance"],
            ["3", "Move to the next flow time and inject decreasing source noise", "Indirect dependence through the refined destination"],
        ],
        [0.55, 3.65, 2.65],
        font_size=8.9,
    )
    add_equation(doc, "flow-path.png", 6.1)
    add_equation(doc, "gaussian-posterior.png", 6.75)
    doc.add_paragraph(
        "A smaller assumed sigma increases the coefficient of the data term and forces closer agreement with a noisy measurement. A larger "
        "assumed sigma weakens data consistency and gives the flow prior more influence. Ill-posed operators can amplify this tradeoff "
        "because the measurement does not constrain every image direction equally."
    )

    doc.add_heading("Mathematical justification", level=1)
    doc.add_heading("Optimal velocity gives the conditional destination", level=2)
    doc.add_paragraph(
        "For conditional flow matching with a straight path, the squared-error optimum is the conditional mean of the target velocity. "
        "Because X1 equals Xt plus one minus t times X1 minus X0, taking the conditional expectation yields Proposition 1 of the paper."
    )
    add_equation(doc, "destination-proof.png", 6.65)
    doc.add_paragraph(
        "This short proof explains why the first step acts as a denoiser: an optimally trained velocity network predicts the conditional "
        "mean destination associated with the current noisy flow state."
    )

    doc.add_heading("Gaussian multiplication gives the refinement", level=2)
    doc.add_paragraph(
        "FLOWER approximates the destination distribution by a Gaussian centered at the prior prediction and combines it with the Gaussian "
        "linear likelihood. Multiplying the two densities adds their quadratic exponents. Collecting the terms in x gives posterior precision "
        "nu to the power minus two times the identity plus sigma to the power minus two times H transpose H. Completing the square gives the "
        "mean and covariance displayed above, which is Proposition 2 in the paper."
    )
    add_equation(doc, "proximal-objective.png", 5.9)
    doc.add_paragraph(
        "Differentiating this objective and setting the gradient to zero produces the same normal equation as the Gaussian mean. This proves "
        "the equivalence between the Bayesian conditional mean and FLOWER's proximal measurement step. It also identifies the exact location "
        "of the noise mismatch: the data-fidelity weight is the reciprocal of the assumed variance."
    )

    doc.add_heading("Progression preserves the conditional path", level=2)
    add_equation(doc, "progression.png", 5.2)
    doc.add_paragraph(
        "Under the paper's independence and Gaussian approximations, the weighted sum of a refined Gaussian destination and independent "
        "standard Gaussian source noise is Gaussian with the mean and covariance stated in Proposition 3. Theorem 1 then follows from the "
        "law of total probability: if the current state has the correct conditional distribution and the transition is correct, integrating "
        "the transition over the current conditional state yields the next conditional distribution."
    )
    add_equation(doc, "ancestral.png", 5.7)
    doc.add_paragraph(
        "Repeating this argument gives an inductive ancestral-sampling interpretation. The paper notes that gamma equal to one is required for "
        "the full covariance sampling argument, while gamma equal to zero gives better reconstruction metrics in practice. This project follows "
        "the reported Flower1-OT reconstruction setting with gamma equal to zero."
    )

    doc.add_heading("Experimental design", level=1)
    add_table(
        doc,
        ["Component", "Protocol"],
        [
            ["Images", "20 fixed random images from each paper-ordered 100-image test pool, seed 2026"],
            ["Pairing", "One measurement per image and task; identical measurement and source-noise trajectory across ratios"],
            ["Control", "Assumed sigma divided by true sigma equals 1.0"],
            ["Model", "Official pretrained optimal-transport flow model for each dataset"],
            ["Metrics", "PSNR and SSIM are higher-is-better; LPIPS is lower-is-better"],
            ["Uncertainty", "95 percent confidence interval across 20 paired image-level differences"],
        ],
        [1.45, 5.4],
        font_size=9.1,
    )
    add_table(
        doc,
        ["Task", "True sigma", "CelebA operator", "AFHQ-Cat operator"],
        [
            ["Denoising", "0.20", "Identity", "Identity"],
            ["Deblurring", "0.05", "61 x 61 Gaussian, blur sigma 1", "61 x 61 Gaussian, blur sigma 3"],
            ["Super-resolution", "0.05", "2x decimation", "4x decimation"],
            ["Random inpainting", "0.01", "70 percent missing", "70 percent missing"],
            ["Box inpainting", "0.05", "Centered 40 x 40 box", "Centered 80 x 80 box"],
        ],
        [1.55, 0.8, 2.25, 2.25],
        font_size=8.6,
    )
    doc.add_paragraph(
        "CelebA uses 100 FLOWER steps for every task. AFHQ-Cat uses 100 steps for denoising and box inpainting, 200 for random inpainting, "
        "and 500 for deblurring and super-resolution, matching the public Flower1-OT scripts. The selected images remain fixed across all tasks."
    )
    add_equation(doc, "delta-psnr.png", 5.9)

    doc.add_page_break()
    doc.add_heading("Results", level=1)
    add_picture(
        doc,
        FIGURES / "figure1_delta_psnr.png",
        6.85,
        "Figure 1  Mean paired PSNR change relative to the correct noise scale. Error bars show 95 percent confidence intervals across 20 images.",
    )
    doc.add_paragraph(
        f"The strongest average response is {DATASET_LABELS[selected_dataset]} {TASK_LABELS[selected_task].lower()} with mean absolute paired "
        f"change {setting_scores[most_sensitive]:.2f} dB across the four mismatched ratios. Its largest tested deviation occurs at ratio "
        f"{worst_ratio:g}, where the mean paired change is {float(worst_row['psnr_delta_vs_correct_mean']):+.2f} dB and "
        f"{loss_fraction:.0%} of the images have lower PSNR than their own control reconstruction. The least responsive setting is "
        f"{DATASET_LABELS[least_sensitive[0]]} {TASK_LABELS[least_sensitive[1]].lower()} at {setting_scores[least_sensitive]:.2f} dB."
    )

    condensed_rows: list[list[str]] = []
    for dataset in DATASET_LABELS:
        for task in TASK_LABELS:
            control = lookup[(dataset, task, 1.0)]
            mismatches = [lookup[(dataset, task, ratio)] for ratio in (0.5, 0.75, 1.25, 1.5)]
            largest = max(mismatches, key=lambda row: abs(float(row["psnr_delta_vs_correct_mean"])))
            paper_psnr = float(config["paper_flower1_ot_reference"][dataset][task][0])
            condensed_rows.append(
                [
                    DATASET_LABELS[dataset],
                    TASK_LABELS[task],
                    f"{float(control['psnr_mean']):.2f}",
                    f"{float(largest['assumed_to_true_sigma']):g}",
                    f"{float(largest['psnr_delta_vs_correct_mean']):+.2f}",
                    f"{paper_psnr:.2f}",
                ]
            )
    add_table(
        doc,
        ["Dataset", "Task", "Control PSNR", "Largest response ratio", "Paired delta dB", "Paper PSNR"],
        condensed_rows,
        [0.85, 1.55, 1.0, 1.05, 1.0, 0.9],
        font_size=8.0,
    )
    doc.add_paragraph(
        "Table 1  Condensed results. The paper values use 100 images, while this project uses a fixed 20-image subset from the same benchmark "
        "pool. The paper column is therefore contextual and not an exact sample-matched replication target."
    ).style = "Caption"

    doc.add_heading("Final reconstruction evidence", level=2)
    add_picture(
        doc,
        FIGURES / "figure2_celeba_qualitative.png",
        6.85,
        "Figure 2  CelebA final reconstructions for the most and least sensitive tasks under the five assumed noise scales.",
    )
    add_picture(
        doc,
        FIGURES / "figure3_afhq_qualitative.png",
        6.85,
        "Figure 3  AFHQ-Cat final reconstructions using the same predeclared selection rule.",
    )
    doc.add_paragraph(
        "The qualitative rows use one image chosen before results were examined. Within each dataset, the plotted tasks are selected by the "
        "highest and lowest mean absolute paired PSNR sensitivity. This avoids choosing examples only because their failures look dramatic. "
        "The images show whether numerical changes correspond to visible residual noise, excessive smoothing, missing detail, or structural artifacts."
    )

    doc.add_heading("Trajectory evidence", level=2)
    add_picture(
        doc,
        FIGURES / "figure4_trajectory.png",
        6.55,
        "Figure 4  FLOWER states at five flow times for the most sensitive dataset-task setting and an image closest to median sensitivity.",
    )
    doc.add_paragraph(
        "Early columns show whether the trajectories begin to separate immediately after measurement refinement. Later columns show whether "
        "the prior compensates for an initial mismatch or whether the discrepancy accumulates. Because the source-noise sequence is shared, "
        "visible separation can be attributed to the assumed likelihood scale and its downstream effect on the state."
    )

    doc.add_heading("Mechanistic evidence", level=2)
    add_picture(
        doc,
        FIGURES / "figure5_internal_states.png",
        6.85,
        "Figure 5  Prior destination, measurement-refined destination, and next trajectory state for three ratios and three flow times.",
    )
    doc.add_paragraph(
        "Sigma does not appear explicitly in the velocity-network destination formula, but previous refinements change the state presented to "
        "the network. The direct effect appears in the gap between the prior destination and the refined destination. Underestimation generally "
        "produces a stronger measurement correction, while overestimation preserves more of the prior prediction. The next state carries that "
        "difference into later iterations."
    )

    doc.add_heading("Conclusions", level=1)
    ranked = sorted(setting_scores, key=setting_scores.get, reverse=True)
    for rank, (dataset, task) in enumerate(ranked[:3], start=1):
        add_bullet(
            doc,
            f"Sensitivity rank {rank}: {DATASET_LABELS[dataset]} {TASK_LABELS[task].lower()} has mean absolute paired PSNR response "
            f"{setting_scores[(dataset, task)]:.2f} dB.",
        )
    under_deltas = [
        float(row["psnr_delta_vs_correct_mean"])
        for row in summary
        if float(row["assumed_to_true_sigma"]) < 1
    ]
    over_deltas = [
        float(row["psnr_delta_vs_correct_mean"])
        for row in summary
        if float(row["assumed_to_true_sigma"]) > 1
    ]
    doc.add_paragraph(
        f"Across all dataset-task settings, the mean PSNR change over the two underestimated ratios is {statistics.mean(under_deltas):+.2f} dB, "
        f"compared with {statistics.mean(over_deltas):+.2f} dB over the two overestimated ratios. This aggregate describes directionality but "
        "does not replace the task-level curves, because inverse operators constrain different image subspaces."
    )
    doc.add_paragraph(
        "The project supports a practical conclusion: correct noise calibration matters most for the settings with the steepest paired curves. "
        "A flat curve indicates tolerance within the tested range, while an asymmetric curve identifies whether conservative or aggressive noise "
        "assumptions are more damaging. These claims apply to the released Flower1-OT models, the tested noise levels, and the fixed 20-image samples."
    )

    doc.add_heading("Limitations", level=1)
    add_bullet(doc, "Twenty images reduce compute cost but give wider confidence intervals than the paper's 100-image tables.")
    add_bullet(doc, "The project tests isotropic Gaussian noise-scale mismatch and does not test an incorrect covariance shape.")
    add_bullet(doc, "The five ratios describe a response curve over a limited range and do not identify behavior under extreme misspecification.")
    add_bullet(doc, "The control uses the correct supplied sigma, but finite model and discretization error can still make another ratio score better on a small sample.")
    add_bullet(doc, "The paper's theoretical posterior argument uses approximations and gamma equal to one, whereas the reconstruction benchmark uses gamma equal to zero for image quality.")

    doc.add_heading("Reproducibility", level=1)
    doc.add_paragraph(
        "The repository records the random image selection, measurements, solver seeds, masks, task parameters, paper step counts, raw metrics, "
        "paired deltas, confidence intervals, selection rules for qualitative figures, and all five figure builders. The Colab notebook stores "
        "results in Google Drive and resumes completed cases after a disconnect."
    )
    add_table(
        doc,
        ["Artifact", "Location"],
        [
            ["Protocol", "config/robustness.json"],
            ["Selected images", "data_splits/*random20_seed2026*"],
            ["Raw results", "artifacts/experiments/main/metrics.csv"],
            ["Paired results", "artifacts/experiments/main/metrics_with_deltas.csv"],
            ["Aggregate results", "artifacts/experiments/main/summary.csv"],
            ["Figures", "artifacts/figures/figure1 through figure5"],
            ["GPU workflow", "notebooks/run_full_benchmark_colab.ipynb"],
        ],
        [1.55, 5.3],
        font_size=8.8,
    )

    doc.add_page_break()
    doc.add_heading("Appendix full results table", level=1)
    full_rows = []
    for row in summary:
        full_rows.append(
            [
                DATASET_LABELS[row["dataset"]],
                TASK_LABELS[row["task"]],
                f"{float(row['assumed_to_true_sigma']):g}",
                f"{float(row['assumed_to_true_variance']):g}",
                metric_with_ci(row, "psnr", 2),
                metric_with_ci(row, "ssim", 3),
                metric_with_ci(row, "lpips", 3),
            ]
        )
    add_table(
        doc,
        ["Dataset", "Task", "Sigma ratio", "Variance ratio", "PSNR", "SSIM", "LPIPS"],
        full_rows,
        [0.72, 1.28, 0.72, 0.78, 1.18, 1.08, 1.08],
        font_size=7.2,
    )
    doc.add_paragraph(
        "Table A1  Mean +/- 95 percent confidence interval across 20 images for every dataset, task, and assumed noise ratio."
    ).style = "Caption"

    doc.add_heading("References", level=1)
    add_reference(doc, 1, "M. Pourya, B. El Rawas, and M. Unser, FLOWER A Flow Matching Solver for Inverse Problems, ICLR 2026. https://openreview.net/forum?id=QGd34p02mI")
    add_reference(doc, 2, "M. Pourya, B. El Rawas, and M. Unser, official FLOWER implementation. https://github.com/mehrsapo/Flower")
    add_reference(doc, 3, "R. Zhang, P. Isola, A. A. Efros, E. Shechtman, and O. Wang, The Unreasonable Effectiveness of Deep Features as a Perceptual Metric, CVPR 2018.")
    add_reference(doc, 4, "Z. Wang, A. C. Bovik, H. R. Sheikh, and E. P. Simoncelli, Image Quality Assessment From Error Visibility to Structural Similarity, IEEE Transactions on Image Processing, 2004.")
    add_reference(doc, 5, "Z. Liu et al., Deep Learning Face Attributes in the Wild, ICCV 2015.")
    add_reference(doc, 6, "Y. Choi et al., StarGAN v2 Diverse Image Synthesis for Multiple Domains, CVPR 2020.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "FLOWER Under Measurement Noise Scale Mismatch"
    doc.core_properties.subject = "Paired robustness study across two datasets and five inverse problems"
    doc.core_properties.author = "Itay Shorian"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
